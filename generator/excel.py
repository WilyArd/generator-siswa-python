"""
Modul penanganan manipulasi file Excel (openpyxl), styling sel, dan pembaruan Data Validation.
"""

import os
import sys
import re
import datetime
import random

try:
    import openpyxl
    from openpyxl.styles import Font, Border, Side
except ImportError:
    openpyxl = None

from generator.jenjang import JENJANG_CONFIG, normalize_jenjang
from generator.student import generate_single_student


def generate_siswa_excel(
    template_path: str,
    output_path: str,
    count: int = 1000,
    level: str = "SMA",
    mode: str = "standard",
    rombel: str = None,
    seed: int = None
) -> str:
    """
    Membaca file template Excel, meng-generate N data siswa, dan menyimpan hasilnya.

    Args:
        template_path: Path ke file master 'format sidigs murid-dapodik.xlsx'
        output_path: Path file Excel hasil generate
        count: Jumlah siswa yang ingin digenerate (default: 1000)
        level: Jenjang sekolah (SMA, SMK, SMP, SD, TK, MA, MTS, MI, SLB)
        mode: Mode kolom ('standard' atau 'full')
        rombel: String kelas/rombel yang ditentukan pengguna (opsional)
        seed: Random seed untuk reproduktibilitas hasil (opsional)

    Returns:
        Path file output yang berhasil disimpan
    """
    if openpyxl is None:
        raise ImportError("Package 'openpyxl' belum terpasang. Jalankan: pip install openpyxl")

    if seed is not None:
        random.seed(seed)

    print(f"\n[+] Memuat template: {template_path}")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"File template '{template_path}' tidak ditemukan!")

    wb = openpyxl.load_workbook(template_path)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    cfg = JENJANG_CONFIG.get(normalize_jenjang(level), JENJANG_CONFIG["SMA"])
    print(f"[+] Sheet aktif: '{sheet_name}'")
    rombel_info = f", Kelas/Rombel: {rombel}" if rombel else (", Kelas: Acak Full" if mode == "full" else ", Kelas: Kosong (Opsional)")
    print(f"[+] Menghasilkan {count:,} data siswa (Jenjang: {cfg['nama']}, Mode: {mode.upper()}{rombel_info})...")

    # Siapkan styling standar
    font_standard = Font(name="Calibri", size=11, bold=False)
    thin_side = Side(style='thin', color='CCCCCC')
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    used_usernames = set()
    start_row = 7
    total_rows = start_row + count - 1

    date_format_str = r'[$-421]d\ mmmm\ yyyy'

    # Setup milestone progress bar
    milestone = max(1, count // 20)

    for i in range(1, count + 1):
        row_num = start_row + i - 1
        student_data = generate_single_student(i, level, mode, used_usernames, rombel_input=rombel)

        for col_idx, val in enumerate(student_data, start=1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = font_standard
            cell.border = cell_border

            # Penanganan tipe data khusus
            if isinstance(val, (datetime.datetime, datetime.date)):
                cell.value = val
                cell.number_format = date_format_str
            elif col_idx in [6, 8, 11, 23, 35, 41, 47, 66]:  # NIS, NISN, NIK, HP, No. KK
                cell.value = str(val) if val is not None else None
                cell.number_format = '@'
            elif isinstance(val, int):
                cell.value = val
                cell.number_format = '0'
            else:
                cell.value = val

        if i % milestone == 0 or i == count:
            pct = (i / count) * 100
            bar = "#" * int(pct // 5) + "-" * (20 - int(pct // 5))
            sys.stdout.write(f"\r    [{bar}] {pct:5.1f}% ({i}/{count} siswa)")
            sys.stdout.flush()

    print("\n[+] Memperbarui aturan Data Validation (Dropdown Excel) sesuai jumlah baris baru...")
    # Update jangkauan Data Validation agar dropdown aktif hingga baris terakhir
    for dv in ws.data_validations.dataValidation:
        sqref_str = str(dv.sqref)
        new_sqref = re.sub(r'1008\b', str(total_rows), sqref_str)
        dv.sqref = new_sqref

    print(f"[+] Menyimpan workbook ke: {output_path} ...")
    wb.save(output_path)
    file_size_mb = os.path.getsize(output_path) / (1024 * 1024)
    print(f"[OK] Sukses! File berhasil dibuat ({file_size_mb:.2f} MB).")
    print(f"     Total Data Siswa: {count:,} siswa (Baris {start_row} s.d {total_rows})")

    return output_path
